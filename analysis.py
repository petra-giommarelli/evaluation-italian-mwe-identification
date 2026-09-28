import pathlib

import matplotlib
import numpy as np
import pandas as pd
from scipy.stats import kruskal, spearmanr

matplotlib.use("Agg")
import matplotlib.pyplot as plt

DATA_DIR = pathlib.Path("data")
OUT_DIR = pathlib.Path("output")

DATASET_FILE = DATA_DIR / "dataset.xlsx"
FREQUENCY_FILE = DATA_DIR / "mwe_frequency.csv"
INFLECTION_FILE = DATA_DIR / "inflection_annotation.xlsx"
SCORES_FILE = DATA_DIR / "model_outputs_scored.xlsx"

MODEL_SHEETS = {"LLaMA": "llama", "Mistral": "mistral", "Gemma": "gemma"}
MODELS = list(MODEL_SHEETS)

OUT_XLSX = OUT_DIR / "analysis_results.xlsx"
OUT_FREQUENCY_FIG = OUT_DIR / "frequency.png"
OUT_CATEGORY_FIG = OUT_DIR / "category.png"
OUT_LENGTH_FIG = OUT_DIR / "length.png"

STATE_MAP = {
    "canonica": "canonical",
    "inflessa contigua": "contiguous-inflected",
    "discontinua": "discontinuous",
}
STATE_ORDER = ["canonical", "contiguous-inflected", "discontinuous"]
CATEGORY_ORDER = ["VMWE", "NMWE", "AMWE", "FMWE"]

GRAYS = {"LLaMA": "#111111", "Mistral": "#B0B0B0", "Gemma": "#6E6E6E"}
LENGTH_EDGES = [0, 10, 20, 30, 40, 50, 60, np.inf]
LENGTH_LABELS = ["3-10", "11-20", "21-30", "31-40", "41-50", "51-60", "60+"]


def normalise(value):
    return str(value).strip().lower().replace("’", "'")


def load_dataset():
    data = pd.read_excel(DATASET_FILE)[["ID", "MWE", "Sentence"]]
    data["key"] = data["MWE"].map(normalise)
    data["n_words"] = data["Sentence"].str.split().str.len()
    return data


def load_frequency():
    frequency = pd.read_csv(FREQUENCY_FILE)
    frequency["key"] = frequency["Espressione"].map(normalise)
    frequency = frequency.rename(
        columns={
            "Categoria": "category",
            "freq_assoluta": "freq_absolute",
            "freq_per_milione": "freq_per_million",
        }
    )
    return frequency[["key", "category", "freq_absolute", "freq_per_million"]].drop_duplicates("key")


def load_model(sheet, valid_ids):
    data = pd.read_excel(SCORES_FILE, sheet_name=sheet)[["ID", "score"]]
    data["score"] = pd.to_numeric(data["score"], errors="coerce")
    data = data[data["ID"].isin(valid_ids)].dropna(subset=["score"])
    return data.drop_duplicates("ID")


def load_inflection(valid_ids):
    data = pd.read_excel(INFLECTION_FILE)[["ID", "stato"]]
    data["state"] = data["stato"].astype(str).str.strip().str.lower().map(STATE_MAP)
    data = data[data["ID"].isin(valid_ids)]
    return data[["ID", "state"]].drop_duplicates("ID")


def state_table(occurrences):
    table = occurrences.pivot_table(index="state", values=MODELS, aggfunc="mean")
    table = table.reindex(STATE_ORDER)[MODELS]
    table["Mean (3 models)"] = table.mean(axis=1)
    long = occurrences.melt(id_vars="state", value_vars=MODELS, value_name="score")
    long = long.dropna(subset=["score"])
    table["% score 3"] = long.groupby("state")["score"].apply(lambda s: (s == 3).mean() * 100).reindex(STATE_ORDER)
    table["% score 0"] = long.groupby("state")["score"].apply(lambda s: (s == 0).mean() * 100).reindex(STATE_ORDER)
    table["n occurrences"] = occurrences.groupby("state").size().reindex(STATE_ORDER)
    return table.round(2)


def format_workbook(path):
    from openpyxl import load_workbook
    from openpyxl.styles import Font

    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        for cell in sheet[1]:
            cell.font = Font(name="Arial", bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = Font(name="Arial")
        for column in sheet.columns:
            width = max((len(str(c.value)) for c in column if c.value is not None), default=10)
            sheet.column_dimensions[column[0].column_letter].width = min(width + 2, 40)
    workbook.save(path)


def main():
    OUT_DIR.mkdir(exist_ok=True)

    dataset = load_dataset()
    valid_ids = set(dataset["ID"])
    n_sentences = len(dataset)
    frequency = load_frequency()

    scores = dataset.merge(frequency, on="key", how="left")
    for name, sheet in MODEL_SHEETS.items():
        model = load_model(sheet, valid_ids).rename(columns={"score": name})
        scores = scores.merge(model, on="ID", how="left")
        missing = scores[name].isna().sum()
        if missing:
            raise SystemExit(f"{name}: {missing} sentences without a score")
    scores["mean_score"] = scores[MODELS].mean(axis=1)

    unmatched = scores["category"].isna().sum()
    if unmatched:
        raise SystemExit(f"{unmatched} sentences whose expression has no frequency entry")

    n_expressions = scores["key"].nunique()
    print(f"Sentences: {n_sentences} | expressions: {n_expressions}")

    dataset_overview = pd.DataFrame(
        [
            {"item": "sentences", "value": n_sentences},
            {"item": "expressions", "value": n_expressions},
            {"item": "models", "value": len(MODELS)},
        ]
    )

    sentence_length = pd.DataFrame(
        [
            {
                "scope": "all",
                "n": n_sentences,
                "mean": round(dataset["n_words"].mean(), 1),
                "median": int(dataset["n_words"].median()),
                "min": int(dataset["n_words"].min()),
                "max": int(dataset["n_words"].max()),
            }
        ]
    )

    length_by_category = (
        scores.groupby("category")["n_words"]
        .agg(["size", "mean", "median", "min", "max"])
        .reindex(CATEGORY_ORDER)
        .rename(columns={"size": "n"})
        .round(1)
    )

    category_counts = (
        scores.drop_duplicates("key")["category"].value_counts().reindex(CATEGORY_ORDER).rename("n_expressions")
    )

    overall_rows = []
    for name in MODELS:
        series = scores[name]
        row = {"Model": name}
        for value in [0, 1, 2, 3]:
            row[f"Score {value}"] = round(100 * (series == value).sum() / len(series), 2)
        row["Mean"] = round(series.mean(), 2)
        overall_rows.append(row)
    overall = pd.DataFrame(overall_rows)

    per_expression = scores.groupby(["key", "category", "freq_absolute", "freq_per_million"])[MODELS].mean()
    per_expression = per_expression.reset_index()
    per_expression["mean_score"] = per_expression[MODELS].mean(axis=1)

    frequency_rows = []
    for label in MODELS + ["mean_score"]:
        rho, p_value = spearmanr(per_expression["freq_per_million"], per_expression[label])
        frequency_rows.append(
            {
                "Model": "Mean (3 models)" if label == "mean_score" else label,
                "Spearman rho": round(rho, 3),
                "p-value": f"{p_value:.2e}",
                "significant (p<0.05)": "yes" if p_value < 0.05 else "no",
            }
        )
    frequency_correlations = pd.DataFrame(frequency_rows)

    category_score3 = pd.DataFrame(
        {name: scores.groupby("category")[name].apply(lambda s: round((s == 3).mean() * 100, 1)) for name in MODELS}
    ).reindex(CATEGORY_ORDER)
    category_score0 = pd.DataFrame(
        {name: scores.groupby("category")[name].apply(lambda s: round((s == 0).mean() * 100, 1)) for name in MODELS}
    ).reindex(CATEGORY_ORDER)

    length_rows = []
    for name in MODELS:
        rho, p_value = spearmanr(scores["n_words"], scores[name])
        length_rows.append({"Model": name, "Spearman rho": round(rho, 3), "p-value": f"{p_value:.2e}"})
    length_correlations = pd.DataFrame(length_rows)

    scores["length_band"] = pd.cut(
        scores["n_words"], bins=LENGTH_EDGES, labels=LENGTH_LABELS, right=True, include_lowest=True
    )
    length_bands = scores.groupby("length_band", observed=False)[MODELS].mean().reindex(LENGTH_LABELS).round(2)
    length_bands["n sentences"] = scores.groupby("length_band", observed=False).size().reindex(LENGTH_LABELS)

    inflection = load_inflection(valid_ids)
    occurrences = scores.merge(inflection, on="ID", how="left")
    annotated = occurrences.dropna(subset=["state"])

    inflection_by_category = pd.crosstab(annotated["category"], annotated["state"], normalize="index")
    inflection_by_category = (inflection_by_category[STATE_ORDER].reindex(CATEGORY_ORDER) * 100).round(1)
    inflection_by_category["% not canonical"] = (100 - inflection_by_category["canonical"]).round(1)
    inflection_by_category["n occurrences"] = (
        annotated.groupby("category").size().reindex(CATEGORY_ORDER)
    )

    verbal = annotated[annotated["category"] == "VMWE"]
    inflection_scores = state_table(verbal)

    frequency_median = (
        verbal.groupby("state")["freq_per_million"].median().reindex(STATE_ORDER).round(3).to_frame("median per million")
    )

    statistic, p_value = kruskal(*[verbal.loc[verbal["state"] == s, "mean_score"] for s in STATE_ORDER])
    kruskal_table = pd.DataFrame(
        [
            {
                "Test": "Kruskal-Wallis across inflectional states (verbal MWEs)",
                "Statistic": round(statistic, 3),
                "p-value": f"{p_value:.2e}",
                "n": len(verbal),
            }
        ]
    )

    print("\nTable 1: score distribution (%) and mean, per model")
    print(overall.to_string(index=False))
    print("\nSentence length (words):")
    print(sentence_length.to_string(index=False))
    print("\nInflectional state within verbal MWEs:")
    print(inflection_scores.to_string())
    print(f"\nKruskal-Wallis: H = {statistic:.3f}, p = {p_value:.2e}, n = {len(verbal)}")
    print("\nMean score by length band:")
    print(length_bands.to_string())

    with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
        dataset_overview.to_excel(writer, sheet_name="dataset_overview", index=False)
        sentence_length.to_excel(writer, sheet_name="sentence_length", index=False)
        length_by_category.to_excel(writer, sheet_name="sentence_length_by_category")
        category_counts.to_frame().to_excel(writer, sheet_name="category_counts")
        overall.to_excel(writer, sheet_name="overall_scores", index=False)
        frequency_correlations.to_excel(writer, sheet_name="frequency_correlations", index=False)
        per_expression.sort_values("freq_per_million").to_excel(writer, sheet_name="scores_per_expression", index=False)
        category_score3.to_excel(writer, sheet_name="category_score3")
        category_score0.to_excel(writer, sheet_name="category_score0")
        length_correlations.to_excel(writer, sheet_name="length_correlations", index=False)
        length_bands.to_excel(writer, sheet_name="length_bands")
        inflection_by_category.to_excel(writer, sheet_name="inflection_by_category")
        inflection_scores.to_excel(writer, sheet_name="inflection_scores_verbal")
        frequency_median.to_excel(writer, sheet_name="inflection_frequency_median")
        kruskal_table.to_excel(writer, sheet_name="inflection_kruskal", index=False)

    format_workbook(OUT_XLSX)

    positive = per_expression[per_expression["freq_per_million"] > 0]
    figure, axis = plt.subplots(figsize=(7, 5))
    axis.scatter(
        positive["freq_per_million"], positive["mean_score"], s=14, alpha=0.35, color="#555555", edgecolor="none"
    )
    axis.set_xscale("log")
    axis.set_xlabel("Frequency (itTenTen, per million, log scale)")
    axis.set_ylabel("Mean identification score (0-3)")
    axis.set_ylim(-0.1, 3.1)
    figure.tight_layout()
    figure.savefig(OUT_FREQUENCY_FIG, dpi=200)
    plt.close(figure)

    figure, axis = plt.subplots(figsize=(7, 5))
    positions = np.arange(len(CATEGORY_ORDER))
    width = 0.25
    for index, name in enumerate(MODELS):
        axis.bar(
            positions + (index - 1) * width,
            category_score3[name].values,
            width,
            label=name,
            color=GRAYS[name],
            edgecolor="black",
            linewidth=0.6,
        )
    axis.set_xticks(positions)
    axis.set_xticklabels(CATEGORY_ORDER)
    axis.set_ylabel("% score 3 (full identification)")
    axis.legend(frameon=False)
    figure.tight_layout()
    figure.savefig(OUT_CATEGORY_FIG, dpi=200)
    plt.close(figure)

    figure, axis = plt.subplots(figsize=(7, 5))
    for name in MODELS:
        axis.plot(
            range(len(LENGTH_LABELS)),
            length_bands[name].values,
            marker="o",
            lw=2,
            label=name,
            color=GRAYS[name],
            markeredgecolor="black",
        )
    axis.set_xticks(range(len(LENGTH_LABELS)))
    axis.set_xticklabels(LENGTH_LABELS)
    axis.set_xlabel("Sentence length (words)")
    axis.set_ylabel("Mean identification score (0-3)")
    axis.set_ylim(0, 3)
    axis.legend(frameon=False)
    figure.tight_layout()
    figure.savefig(OUT_LENGTH_FIG, dpi=200)
    plt.close(figure)

    print(f"\nSaved {OUT_XLSX}, {OUT_FREQUENCY_FIG}, {OUT_CATEGORY_FIG}, {OUT_LENGTH_FIG}")


if __name__ == "__main__":
    main()